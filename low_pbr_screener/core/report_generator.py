"""Excel 리포트 생성 모듈. 5개 시트로 구성.

전략별/타이밍별 Excel 자동필터 + 쉬운 말 발굴이유/타이밍이유 포함.
"""

from datetime import datetime
from pathlib import Path

import pandas as pd

from .config import CATEGORIES_CONFIG, OUTPUT_DIR
from .screener import generate_simple_reason, generate_timing_explanation

# CATEGORIES_CONFIG 키 순서 그대로 사용 (동적 AI 카테고리 지원)
_CAT_NAMES = list(CATEGORIES_CONFIG.keys())
_CAT_ICONS = [(n, CATEGORIES_CONFIG[n].get("emoji", "")) for n in _CAT_NAMES]

TIMING_KR = {"STRONG_BUY": "즉시매수", "BUY_NOW": "매수", "WATCH": "관심"}
TIMING_RANK = {"STRONG_BUY": 3, "BUY_NOW": 2, "WATCH": 1, "": 0}

# 컬럼 한국어 설명 매핑
_COLUMN_KR = {
    "Name": "종목명",
    "Market": "시장구분",
    "Sector": "업종",
    "Close": "현재가",
    "Low52": "52주최저가",
    "High52": "52주최고가",
    "MarketCap": "시가총액(원)",
    "PBR": "주가/순자산",
    "PER": "주가/순이익",
    "BPS": "주당순자산",
    "EPS": "주당순이익",
    "DIV": "배당수익률(%)",
    "DPS": "주당배당금",
    "ROE": "자기자본이익률",
    "ROA": "총자산이익률",
    "QualityGrade": "우량등급(A/B/C)",
    "QualityTags": "우량태그",
    "QualityWarnings": "경고태그",
    "RevenueGrowth": "매출성장률(%)",
    "RevCAGR3Y": "3년매출CAGR(%)",
    "OperatingMargin": "영업이익률(%)",
    "OpMarginPrev": "전기영업이익률(%)",
    "DebtRatio": "부채비율(%)",
    "OpCashFlow": "영업현금흐름(억)",
    "OpCashFlowPrev": "전기영업현금흐름(억)",
    "ICR": "이자보상배율",
    "NetDebtEBITDA": "순부채/EBITDA",
    "EPSGrowth": "EPS성장률(%)",
    "PEG": "PEG비율",
    "PEGVerdict": "PEG판정",
    "RSI14": "상대강도(14일)",
    "MACD": "MACD값",
    "MACD_Signal": "MACD시그널",
    "MACD_Hist": "MACD히스토그램",
    "StochK": "스토캐스틱K",
    "StochD": "스토캐스틱D",
    "MA20": "20일이동평균",
    "Flow1M_FI_KRW": "외인1개월순매수(원)",
    "Flow1M_INST_KRW": "기관1개월순매수(원)",
    "NetBuyDays_FI_10d": "외인순매수일(10일중)",
    "NetBuyDays_INST_10d": "기관순매수일(10일중)",
    "Near52wLow": "52주저점근접도",
    "DrawdownFrom52wHigh": "고점대비하락률",
    "AvgTradingValue20D": "20일평균거래대금",
    "CashRatio": "현금/시총비율",
    "RealEstateRatio": "부동산/시총비율",
    "operating_profit": "영업이익(억)",
    "시총(조)": "시총(조원)",
    "종합점수": "전략최고점수",
    "매수타이밍": "매수시점판정",
    "통과전략": "통과한전략목록",
    "발굴이유_요약": "왜 선별됐나",
    "타이밍이유_요약": "타이밍근거",
    "매수트리거": "직접입력",
    "목표가": "직접입력",
    "손절가": "직접입력",
    "메모": "직접입력",
}


def _format_number(val, fmt="{:,.0f}"):
    try:
        return fmt.format(float(val))
    except (ValueError, TypeError):
        return str(val)


def _best_timing(row):
    """종목의 최고 타이밍 등급(한글) 반환."""
    best = 0
    for name in _CAT_NAMES:
        t = str(row.get(f"Timing_{name}", ""))
        best = max(best, TIMING_RANK.get(t, 0))
    return {3: "즉시매수", 2: "매수", 1: "관심"}.get(best, "-")


def _passed_strategies(row):
    """통과한 전략 이름들을 쉼표로 연결."""
    passed = []
    for name in _CAT_NAMES:
        if row.get(f"Pass_{name}", False):
            emoji = CATEGORIES_CONFIG[name].get("emoji", "")
            passed.append(f"{emoji}{name}")
    return ", ".join(passed) if passed else "-"


def generate_excel_report(universe_df: pd.DataFrame, cfg: dict, date: str) -> Path:
    """Excel 5시트 리포트 생성.

    시트: README, CONFIG, UNIVERSE, CANDIDATES, DEEPDIVE_TOP20
    """
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"급등후보주스크리닝_{date}_{timestamp}.xlsx"
    filepath = OUTPUT_DIR / filename

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        # 1) README 시트
        readme_items = [
            ("프로그램", "급등후보주 발굴 프로그램 v3.0"),
            ("기준일", date),
            ("생성시각", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("전종목 수", len(universe_df)),
        ]
        for name in _CAT_NAMES:
            key = f"Pass_{name}"
            n = int(universe_df[key].sum()) if key in universe_df.columns else 0
            emoji = CATEGORIES_CONFIG[name].get("emoji", "")
            readme_items.append((f"{emoji} {name} 통과", n))

        if "Pass_any" in universe_df.columns:
            readme_items.append(("전체 통과(합산)", int(universe_df["Pass_any"].sum())))
        if "QualityGrade" in universe_df.columns:
            n_c = int((universe_df["QualityGrade"] == "C").sum())
            if n_c > 0:
                readme_items.append(("⚠ C등급 경고", n_c))
        # ── 시트별 설명 ──
        readme_items.append(("", ""))
        readme_items.append(("═══ 시트 안내 ═══", ""))
        readme_items.append(
            (
                "① CONFIG",
                "스크리닝에 사용된 전략별 게이트/스코어링 파라미터 전체 목록. "
                "어떤 조건으로 필터링했는지 설정값을 확인할 수 있다.",
            )
        )
        readme_items.append(
            (
                "② UNIVERSE",
                "시장 전종목(PBR 범위 내) 원본 데이터. "
                "가격·밸류에이션·재무·기술·수급 지표 + 전략별 Pass/Score/Timing이 모두 포함된다. "
                "필터를 직접 걸어 나만의 조건으로 탐색할 때 사용.",
            )
        )
        readme_items.append(
            (
                "③ CANDIDATES",
                "1개 이상 전략을 통과한 종목만 추출한 핵심 시트. "
                "종합점수 내림차순 정렬. '매수타이밍'(즉시매수/매수/관심)과 '통과전략' 컬럼으로 "
                "자동필터하면 빠르게 후보를 좁힐 수 있다. "
                "'발굴이유_요약'과 '타이밍이유_요약'에 쉬운 말 설명 포함.",
            )
        )
        readme_items.append(
            (
                "④ DEEPDIVE_TOP20",
                "CANDIDATES 상위 20종목을 깊이 분석용으로 뽑은 시트. "
                "모든 재무·기술·수급 컬럼 + 직접 입력란(매수트리거/목표가/손절가/메모) 제공. "
                "실전 매매 전 최종 검토·기록에 활용.",
            )
        )
        readme_items.append(("", ""))
        readme_items.append(
            (
                "💡 활용 팁",
                "CANDIDATES 시트에서 '매수타이밍=즉시매수'로 필터 → 종합점수 상위 종목을 "
                "DEEPDIVE에서 재무·차트 검토 → 매수트리거/목표가 기입 후 실행",
            )
        )

        readme_data = {
            "항목": [r[0] for r in readme_items],
            "내용": [str(r[1]) for r in readme_items],
        }
        pd.DataFrame(readme_data).to_excel(writer, sheet_name="README", index=False)

        # 2) CONFIG 시트
        config_rows = []
        _flatten_config(cfg, "", config_rows)
        pd.DataFrame(config_rows, columns=["파라미터", "값"]).to_excel(
            writer, sheet_name="CONFIG", index=False
        )

        # 3) UNIVERSE 시트
        universe_cols = _get_universe_columns(universe_df)
        universe_df[universe_cols].to_excel(writer, sheet_name="UNIVERSE", index=True)

        # 4) CANDIDATES 시트 — 필터용 컬럼 추가
        pass_any_key = "Pass_any"
        if pass_any_key in universe_df.columns:
            candidates = universe_df[universe_df[pass_any_key]].copy()
        else:
            pass_cols = [
                f"Pass_{n}" for n in _CAT_NAMES if f"Pass_{n}" in universe_df.columns
            ]
            if pass_cols:
                candidates = universe_df[universe_df[pass_cols].any(axis=1)].copy()
            else:
                candidates = universe_df.head(0).copy()

        # 시가총액(조) 컬럼: 1조=1.0, 1000억=0.1
        if "MarketCap" in candidates.columns:
            candidates["시총(조)"] = (
                candidates["MarketCap"].fillna(0).astype(float) / 1e12
            ).round(2)

        # 쉬운 말 컬럼 생성
        candidates["매수타이밍"] = candidates.apply(_best_timing, axis=1)
        candidates["통과전략"] = candidates.apply(_passed_strategies, axis=1)
        candidates["발굴이유_요약"] = candidates.apply(generate_simple_reason, axis=1)
        candidates["타이밍이유_요약"] = candidates.apply(
            lambda row: generate_timing_explanation(row, _CAT_ICONS), axis=1
        )

        # 종합점수 (전략별 최고 점수) + 정렬
        score_cols = [
            f"Score_{n}" for n in _CAT_NAMES if f"Score_{n}" in candidates.columns
        ]
        if score_cols:
            candidates["종합점수"] = candidates[score_cols].max(axis=1)
            candidates = candidates.sort_values("종합점수", ascending=False)
        else:
            candidates["종합점수"] = 0

        cand_cols = _get_candidate_columns(candidates)
        candidates[cand_cols].to_excel(writer, sheet_name="CANDIDATES", index=True)

        # 5) DEEPDIVE_TOP20 시트
        top20 = candidates.head(20).copy()
        top20["매수트리거"] = ""
        top20["목표가"] = ""
        top20["손절가"] = ""
        top20["메모"] = ""
        deep_cols = _get_deepdive_columns(top20)
        top20[deep_cols].to_excel(writer, sheet_name="DEEPDIVE_TOP20", index=True)

        # ── 한국어 설명 행 삽입 + 자동필터 적용 ──
        for sheet_name in ["UNIVERSE", "CANDIDATES", "DEEPDIVE_TOP20"]:
            ws = writer.sheets[sheet_name]
            _insert_kr_desc_row(ws)
            ws.auto_filter.ref = ws.dimensions

    return filepath


def _get_col_kr(col_name: str) -> str:
    """컬럼명에 대응하는 한국어 설명 반환."""
    if col_name in _COLUMN_KR:
        return _COLUMN_KR[col_name]
    # 동적 전략별 컬럼: Pass_전략, Score_전략, Timing_전략, Reason_전략
    for prefix, kr in [
        ("Pass_", "통과여부"),
        ("Score_", "점수"),
        ("Timing_", "타이밍"),
        ("Reason_", "근거"),
    ]:
        if col_name.startswith(prefix):
            return kr
    return ""


def _insert_kr_desc_row(ws):
    """openpyxl 워크시트의 헤더(1행) 아래에 한국어 설명 행을 삽입."""
    from openpyxl.styles import Alignment, Font, PatternFill

    ws.insert_rows(2)
    fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    font = Font(size=9, color="666666")
    align = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, ws.max_column + 1):
        header_cell = ws.cell(row=1, column=col_idx)
        header_val = str(header_cell.value) if header_cell.value else ""
        kr = _get_col_kr(header_val)
        desc_cell = ws.cell(row=2, column=col_idx, value=kr)
        desc_cell.fill = fill
        desc_cell.font = font
        desc_cell.alignment = align


def _flatten_config(obj, prefix, rows):
    if isinstance(obj, dict):
        for k, v in obj.items():
            new_prefix = f"{prefix}.{k}" if prefix else k
            _flatten_config(v, new_prefix, rows)
    elif isinstance(obj, list):
        rows.append((prefix, str(obj)))
    else:
        rows.append((prefix, str(obj)))


def _get_universe_columns(df: pd.DataFrame) -> list[str]:
    desired = [
        "Name",
        "Market",
        "Sector",
        "Close",
        "Low52",
        "High52",
        "MarketCap",
        "PBR",
        "PER",
        "BPS",
        "EPS",
        "DIV",
        "DPS",
        "ROE",
        "ROA",
        "QualityGrade",
        "QualityTags",
        "QualityWarnings",
        "RevenueGrowth",
        "RevCAGR3Y",
        "OperatingMargin",
        "OpMarginPrev",
        "DebtRatio",
        "OpCashFlow",
        "OpCashFlowPrev",
        "ICR",
        "NetDebtEBITDA",
        "EPSGrowth",
        "PEG",
        "PEGVerdict",
        "RSI14",
        "MACD",
        "MACD_Signal",
        "MACD_Hist",
        "StochK",
        "StochD",
        "MA20",
        "Flow1M_FI_KRW",
        "Flow1M_INST_KRW",
        "NetBuyDays_FI_10d",
        "NetBuyDays_INST_10d",
        "Near52wLow",
        "DrawdownFrom52wHigh",
        "AvgTradingValue20D",
        "CashRatio",
        "RealEstateRatio",
        "operating_profit",
    ]
    for name in _CAT_NAMES:
        desired.extend(
            [f"Pass_{name}", f"Score_{name}", f"Timing_{name}", f"Reason_{name}"]
        )
    desired.append("Pass_any")
    return [c for c in desired if c in df.columns]


def _get_candidate_columns(df: pd.DataFrame) -> list[str]:
    # 핵심 지표 → 우량주 → 필터용 요약 컬럼 → 전략별 상세
    desired = [
        "Name",
        "Market",
        "Sector",
        "시총(조)",
        "Close",
        "PBR",
        "PER",
        "EPS",
        "ROE",
        "ROA",
        "DIV",
        "QualityGrade",
        "QualityTags",
        "QualityWarnings",
        "RevenueGrowth",
        "RevCAGR3Y",
        "OperatingMargin",
        "OpMarginPrev",
        "DebtRatio",
        "OpCashFlow",
        "OpCashFlowPrev",
        "ICR",
        "NetDebtEBITDA",
        "EPSGrowth",
        "PEG",
        "PEGVerdict",
        "Low52",
        "High52",
        "CashRatio",
        "RealEstateRatio",
        "operating_profit",
        "RSI14",
        "MACD_Hist",
        "StochK",
        "Flow1M_FI_KRW",
        "Flow1M_INST_KRW",
        # 필터용 요약 컬럼
        "종합점수",
        "매수타이밍",
        "통과전략",
        "발굴이유_요약",
        "타이밍이유_요약",
    ]
    for name in _CAT_NAMES:
        desired.extend([f"Pass_{name}", f"Score_{name}", f"Timing_{name}"])
    return [c for c in desired if c in df.columns]


def _get_deepdive_columns(df: pd.DataFrame) -> list[str]:
    desired = [
        "Name",
        "Market",
        "Sector",
        "시총(조)",
        "Close",
        "Low52",
        "High52",
        "PBR",
        "PER",
        "EPS",
        "ROE",
        "ROA",
        "DIV",
        "DPS",
        "QualityGrade",
        "QualityTags",
        "QualityWarnings",
        "RevenueGrowth",
        "RevCAGR3Y",
        "OperatingMargin",
        "OpMarginPrev",
        "DebtRatio",
        "OpCashFlow",
        "OpCashFlowPrev",
        "ICR",
        "NetDebtEBITDA",
        "EPSGrowth",
        "PEG",
        "PEGVerdict",
        "CashRatio",
        "RealEstateRatio",
        "operating_profit",
        "RSI14",
        "MACD",
        "MACD_Signal",
        "MACD_Hist",
        "StochK",
        "StochD",
        "MA20",
        "Flow1M_FI_KRW",
        "Flow1M_INST_KRW",
        "NetBuyDays_FI_10d",
        "NetBuyDays_INST_10d",
        # 요약
        "종합점수",
        "매수타이밍",
        "통과전략",
        "발굴이유_요약",
        "타이밍이유_요약",
    ]
    for name in _CAT_NAMES:
        desired.extend([f"Score_{name}", f"Timing_{name}"])
    desired.extend(["매수트리거", "목표가", "손절가", "메모"])
    return [c for c in desired if c in df.columns]
